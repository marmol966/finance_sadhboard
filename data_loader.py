# -*- coding: utf-8 -*-
"""
Lectura de finance.xlsx (SOLO LECTURA) y transformacion a estructuras
listas para pintar en el dashboard.

finance.xlsx vive FUERA de esta carpeta (G:\\Mi unidad\\Marmol\\3-Finances\\finance.xlsx).
Este modulo nunca escribe en ese archivo ni en ningun otro fuera de
"Finance Dashboard App".

Nota de rendimiento: las hojas se leen con openpyxl en modo read_only, que
esta optimizado para iteracion SECUENCIAL (iter_rows). Acceder a celdas
sueltas por coordenada (ws["H12"]) en una hoja read_only obliga a openpyxl a
re-recorrer el XML desde el principio cada vez que "saltamos hacia atras",
lo que en una hoja como Portfolio (varias tablas independientes, saltos de
fila constantes) puede tardar minutos. Por eso aqui cada hoja se vuelca UNA
sola vez a una lista en memoria (`_Grid`) y toda la logica de parseo trabaja
sobre esa copia en memoria con acceso O(1).
"""

import datetime
import os
import re

import openpyxl
from openpyxl.utils import column_index_from_string

# finance.xlsx esta un nivel por encima de "Finance Dashboard App"
APP_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.normpath(os.path.join(APP_DIR, "..", "finance.xlsx"))

CARTERA_ORDER = ["Cartera Old", "Cartera New", "RD26", "DEFI", "HYPE"]
RIESGO_ORDER = ["Liquided", "LP", "Medium", "Gem", "Crash"]
USUFRUCTO_ORDER = ["Hold", "Stake", "Farm"]

# Evolucion!X:AJ -> serie semanal de valor actual y capital invertido por
# cada una de las 5 carteras cripto. Ver _load_cartera_evolucion.
#
# Layout verificado en la cabecera Evolucion!fila 2: cada cartera ocupa un
# bloque [actual, invertido, % rendimiento] salvo 'Cartera Old', que solo
# tiene [actual, invertido] (sin columna de % en la hoja). El usuario edita
# esta hoja a mano y puede volver a insertar/mover columnas — si las cifras
# dejan de reconciliar (ver _validate), reverificar esta cabecera antes de
# asumir que estas letras siguen siendo validas.
CARTERA_EVOLUCION_COLS = {
    "Cartera Old": ("X", "Y"),
    "Cartera New": ("Z", "AA"),
    "RD26": ("AF", "AG"),
    "DEFI": ("AC", "AD"),
    "HYPE": ("AI", "AJ"),
}

# Evolucion!AL:AP -> desglose semanal por nivel de riesgo (Liquided/LP/
# Medium/Gem/Crash) del conjunto de las 5 carteras. Ver _load_cripto_riesgo_evolucion.
RIESGO_EVOLUCION_COLS = {
    "Liquided": "AL",
    "LP": "AM",
    "Medium": "AN",
    "Gem": "AO",
    "Crash": "AP",
}

# Evolucion!AQ:AS -> desglose semanal por usufructo (Hold/Stake/Farm). Ver
# _load_cripto_uso_evolucion.
USO_EVOLUCION_COLS = {
    "Hold": "AQ",
    "Stake": "AR",
    "Farm": "AS",
}

# Evolucion!S:W -> desglose semanal por categoria de Renta Variable
# (Liquided/Fondos/MP/ETF/Acciones). Ya viene en EUR (a diferencia de los
# bloques cripto, aqui no hace falta convertir con BD). Ver
# _load_rv_categoria_evolucion.
RV_CATEGORIA_EVOLUCION_COLS = {
    "LIQUIDED": "S",
    "FONDOS": "T",
    "MP": "U",
    "ETF": "V",
    "ACC": "W",
}

# Portfolio!M9:N164 (posiciones cripto individuales) tiene, ademas de M/N
# (riesgo/coin) y O/P/Q (capital/cantidad/precio agregados), una columna de
# CAPITAL propia por cada una de las 5 carteras individuales (no
# documentada con cabecera junto a M:Q — se localizo comprobando que la
# suma de cada columna, para todas las posiciones, coincide EXACTAMENTE con
# el 'Actual' oficial de esa cartera en GUI!D8:D12). Capital en USD sin
# convertir, igual que O/Q. Ver _load_cripto_top_por_cartera.
CARTERA_CAPITAL_COL = {
    "Cartera Old": "X",
    "Cartera New": "BP",
    "RD26": "CI",
    "DEFI": "CL",
    "HYPE": "DE",
}

_COORD_RE = re.compile(r"^([A-Za-z]+)(\d+)$")


class _Grid:
    """Copia en memoria (lista de listas) de una hoja, con acceso tipo A1."""

    def __init__(self, rows):
        self.rows = rows  # tuple de tuplas, 0-indexado

    def get(self, coord):
        m = _COORD_RE.match(coord)
        col_letters, row_str = m.group(1), m.group(2)
        row_idx = int(row_str) - 1
        col_idx = column_index_from_string(col_letters) - 1
        if row_idx >= len(self.rows):
            return None
        row = self.rows[row_idx]
        if col_idx >= len(row):
            return None
        return row[col_idx]


def _iter_fechas(g, start_row, max_blank=4, max_scan=2000):
    """Itera (fila, fecha) de Evolucion!D desde start_row mientras haya
    fechas validas, parando tras `max_blank` filas seguidas sin fecha.

    La hoja Evolucion reutiliza las mismas columnas para mas de una tabla
    verticalmente (p.ej. despues del historico semanal denso hay un hueco en
    blanco y luego un resumen a intervalos CRECIENTES yendo hacia atras en
    el tiempo — no equiespaciado, para otro proposito). Parar por "varias
    filas en blanco seguidas" en vez de por un numero de fila fijo evita
    mezclar ambas tablas y sigue funcionando aunque la hoja crezca semana a
    semana (un limite fijo se quedaria corto a los pocos meses).
    """
    blancos = 0
    for r in range(start_row, start_row + max_scan):
        fecha = g.get(f"D{r}")
        if not isinstance(fecha, datetime.datetime):
            blancos += 1
            if blancos >= max_blank:
                return
            continue
        blancos = 0
        yield r, fecha


def _norm_riesgo(valor):
    if valor is None:
        return None
    v = str(valor).strip()
    if v.upper() == "GEM":
        return "Gem"
    return v


def load_workbook():
    """Abre finance.xlsx en modo solo-lectura (nunca se guarda)."""
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"No se encuentra finance.xlsx en: {EXCEL_PATH}")
    return openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)


def _read_grid(wb, sheet_name, max_row, max_col):
    ws = wb[sheet_name]
    rows = tuple(ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True))
    return _Grid(rows)


def load_data():
    wb = load_workbook()
    # Una sola pasada secuencial por cada hoja que necesitamos (rapido); el
    # resto de la funcion trabaja sobre estas copias en memoria.
    g_cashflow = _read_grid(wb, "CashFlow", max_row=20, max_col=12)
    g_portfolio = _read_grid(wb, "Portfolio", max_row=165, max_col=110)
    g_gui = _read_grid(wb, "GUI", max_row=95, max_col=27)
    g_evolucion = _read_grid(wb, "Evolucion", max_row=450, max_col=58)
    wb.close()

    # Portfolio!D3 = tipo de cambio USD->EUR (los precios/capitales de las
    # posiciones cripto individuales, columnas O/Q, estan en USD; el resto
    # de la hoja -bancos, RV, totales de CashFlow- ya esta en EUR).
    fx_usd_eur = g_portfolio.get("D3") or 1.0

    data = {}
    data["fx_usd_eur"] = fx_usd_eur
    data["patrimonio"] = _load_patrimonio(g_cashflow)
    data["bancos"] = _load_bancos(g_portfolio)
    data["activos_otros"] = _load_activos_otros(g_portfolio)
    data["cripto_cartera"] = _load_cripto_cartera(g_portfolio)
    data["cripto_posiciones"] = _load_cripto_posiciones(g_portfolio, fx_usd_eur)
    data["cripto_top_por_cartera"] = _load_cripto_top_por_cartera(g_portfolio, fx_usd_eur)
    data["cripto_riesgo"] = _aggregate_riesgo(data["cripto_posiciones"])
    data["cripto_cross"] = _load_cripto_cross(g_gui, fx_usd_eur)
    data["cripto_uso_cross"] = _load_cripto_uso_cross(g_gui, fx_usd_eur)
    data["cripto_rendimiento"] = _load_cripto_rendimiento(g_gui, fx_usd_eur)
    data["rv_posiciones"] = _load_rv_posiciones(g_portfolio)
    data["rv_categoria"] = _load_rv_categoria(g_portfolio)
    data["rv_resumen"] = _load_rv_resumen(g_gui)
    data["evolucion"] = _load_evolucion(g_evolucion)
    data["rv_evolucion"] = _load_rv_evolucion(g_evolucion)
    data["rv_categoria_evolucion"] = _load_rv_categoria_evolucion(g_evolucion)

    filas_fecha_cripto = list(_iter_fechas(g_evolucion, start_row=205))
    fx_por_fila = _fx_semanal_por_fila(g_evolucion, filas_fecha_cripto)
    data["cartera_evolucion"] = _load_cartera_evolucion(g_evolucion, filas_fecha_cripto, fx_por_fila)
    data["cripto_riesgo_evolucion"] = _load_cripto_riesgo_evolucion(g_evolucion, filas_fecha_cripto, fx_por_fila)
    data["cripto_uso_evolucion"] = _load_cripto_uso_evolucion(g_evolucion, filas_fecha_cripto, fx_por_fila)

    total = data["patrimonio"]["total"]
    total_cripto = data["patrimonio"]["cripto"]
    total_rv = data["patrimonio"]["renta_variable"]

    for pos in data["cripto_posiciones"]:
        pos["pct_categoria"] = (pos["capital"] / total_cripto * 100) if total_cripto else 0
        pos["pct_patrimonio"] = (pos["capital"] / total * 100) if total else 0

    for pos in data["rv_posiciones"]:
        pos["pct_categoria"] = (pos["valor_actual"] / total_rv * 100) if total_rv else 0
        pos["pct_patrimonio"] = (pos["valor_actual"] / total * 100) if total else 0

    data["warnings"] = _validate(data)
    return data


def _validate(data):
    """Comprueba que las distintas fuentes del Excel reconcilian entre si.

    finance.xlsx vive en una unidad de Google Drive que el usuario edita en
    vivo; una lectura a medio-sincronizar puede devolver una hoja vacia o
    incompleta sin que openpyxl lance ningun error (simplemente hay menos
    filas). Estas comprobaciones existen para detectar esa situacion -no
    para validar reglas de negocio- asi que el umbral es deliberadamente
    laxo (1 EUR) y solo compara cosas que YA deberian cuadrar al centimo.
    """
    warnings = []
    pat = data["patrimonio"]
    tol = 1.0

    suma_categorias = pat["banca"] + pat["activos_otros"] + pat["renta_variable"] + pat["cripto"]
    if abs(suma_categorias - pat["total"]) > tol:
        warnings.append(
            f"CashFlow: BANCA+ACTIVOS+RV+Cripto ({suma_categorias:,.2f}) no coincide "
            f"con PATRIMONIO ({pat['total']:,.2f})."
        )

    suma_bancos = sum(b["importe"] for b in data["bancos"])
    if abs(suma_bancos - pat["banca"]) > tol:
        warnings.append(
            f"Bancos: la suma de posiciones ({suma_bancos:,.2f}) no coincide con "
            f"CashFlow!BANCA ({pat['banca']:,.2f}) — posible lectura incompleta de Portfolio."
        )

    suma_activos = sum(a["importe"] for a in data["activos_otros"])
    if abs(suma_activos - pat["activos_otros"]) > tol:
        warnings.append(
            f"Activos: la suma de posiciones ({suma_activos:,.2f}) no coincide con "
            f"CashFlow!ACTIVOS ({pat['activos_otros']:,.2f})."
        )

    suma_cripto_cartera = sum(x["actual"] for x in data["cripto_cartera"])
    suma_cripto_pos = sum(p["capital"] for p in data["cripto_posiciones"])
    if abs(suma_cripto_cartera - pat["cripto"]) > tol or abs(suma_cripto_pos - pat["cripto"]) > tol:
        warnings.append(
            f"Cripto: cartera={suma_cripto_cartera:,.2f}, posiciones={suma_cripto_pos:,.2f}, "
            f"CashFlow!Cripto={pat['cripto']:,.2f} — no cuadran, probable lectura parcial "
            "del Excel (reintenta actualizar)."
        )
    if pat["cripto"] < tol and suma_cripto_pos > tol:
        warnings.append(
            "Cripto muestra 0€ en el resumen pero hay posiciones con valor: la hoja "
            "CashFlow probablemente se leyo a medias. Pulsa 'Actualizar datos'."
        )

    suma_rv = sum(p["valor_actual"] for p in data["rv_posiciones"])
    suma_rv_cat = sum(x["importe"] for x in data["rv_categoria"])
    if abs(suma_rv - pat["renta_variable"]) > tol or abs(suma_rv_cat - pat["renta_variable"]) > tol:
        warnings.append(
            f"Renta variable: posiciones={suma_rv:,.2f}, por categoria={suma_rv_cat:,.2f}, "
            f"CashFlow!Renta Variable={pat['renta_variable']:,.2f} — no cuadran."
        )
    categorias_desconocidas = sorted({
        p["categoria"] for p in data["rv_posiciones"]
        if p["categoria"] not in ("LIQUIDED", "FONDOS", "MP", "ETF", "ACC")
    })
    if categorias_desconocidas:
        warnings.append(
            f"Renta variable: hay posiciones con categoria no reconocida: {categorias_desconocidas} "
            "(la app solo sabe colorear LIQUIDED/FONDOS/MP/ETF/ACC)."
        )

    rv_resumen = data["rv_resumen"]
    if rv_resumen["actual"] is None or rv_resumen["invertido"] is None:
        warnings.append(
            "Renta variable: no se pudo leer GUI!RV / GUI!RV_Inv (capital invertido total)."
        )
    elif abs(rv_resumen["actual"] - pat["renta_variable"]) > tol:
        warnings.append(
            f"Renta variable: GUI!RV ({rv_resumen['actual']:,.2f}) no coincide con "
            f"CashFlow!Renta Variable ({pat['renta_variable']:,.2f})."
        )

    return warnings


def _load_patrimonio(g):
    """CashFlow!C:D -> resumen de patrimonio, buscado POR ETIQUETA (columna C)
    en vez de por numero de fila fijo. Las filas de este resumen se han
    desplazado alguna vez al anadir/quitar una sub-linea (p.ej. una divisa
    bajo ACTIVOS), asi que anclar por texto es mas robusto que un rango fijo.
    """
    labels = {}
    for r in range(1, 15):
        label = g.get(f"C{r}")
        if label is None:
            continue
        value = g.get(f"D{r}")
        labels[str(label).strip().upper()] = value

    def pick(*keys):
        for k in keys:
            v = labels.get(k)
            if isinstance(v, (int, float)):
                return v
        return 0

    return {
        "total": pick("PATRIMONIO"),
        "banca": pick("BANCA"),
        "activos_otros": pick("ACTIVOS", "ACTIVO"),
        "inversiones": pick("INVERSIONES"),
        "renta_variable": pick("RENTA VARIABLE"),
        "cripto": pick("CRIPTO"),
    }


def _load_bancos(g):
    """Tabla BANCOS: Portfolio!C9:E~18 (Proveedor / Descripcion / Cantidad).

    Se recorre un rango amplio y se descartan filas sin Descripcion+Cantidad
    simultaneas, para no depender de un numero de filas fijo si el usuario
    anade/quita bancos en el futuro.
    """
    filas = []
    for r in range(9, 23):
        proveedor = g.get(f"C{r}")
        desc = g.get(f"D{r}")
        cantidad = g.get(f"E{r}")
        if desc is None or not isinstance(cantidad, (int, float)):
            continue
        filas.append({
            "proveedor": proveedor if proveedor else desc,
            "detalle": desc,
            "importe": cantidad,
        })
    return filas


def _load_activos_otros(g):
    """Tabla ACTIVOS (moto, divisas...): Portfolio!C24:E~28."""
    filas = []
    for r in range(24, 29):
        tipo = g.get(f"C{r}")
        desc = g.get(f"D{r}")
        cantidad = g.get(f"E{r}")
        if desc is None or not isinstance(cantidad, (int, float)):
            continue
        filas.append({"tipo": tipo, "detalle": desc, "importe": cantidad})
    return filas


def _load_cripto_cartera(g):
    """Portfolio!H12:K16 -> valor actual EN EUR por cartera/estrategia.

    Esta tabla ya viene en EUR en el propio Excel y su suma reconcilia
    exactamente con el total de cripto de CashFlow. No se usa GUI!D8:D12
    porque esa columna resulto estar en USD sin convertir (ver README:
    limitaciones conocidas).
    """
    filas = []
    for r in range(9, 20):
        nombre = g.get(f"H{r}")
        total_eur = g.get(f"K{r}")
        if not nombre or not isinstance(total_eur, (int, float)):
            continue
        filas.append({"cartera": nombre, "actual": total_eur})
    filas.sort(key=lambda x: CARTERA_ORDER.index(x["cartera"]) if x["cartera"] in CARTERA_ORDER else 99)
    return filas


def _aggregate_riesgo(posiciones):
    """Agrega las posiciones cripto (ya en EUR) por nivel de riesgo."""
    agg = {}
    for p in posiciones:
        r = p["riesgo"]
        entry = agg.setdefault(r, {"riesgo": r, "actual": 0.0, "n_posiciones": 0})
        entry["actual"] += p["capital"]
        if p["capital"] > 0.01:
            entry["n_posiciones"] += 1
    filas = list(agg.values())
    filas.sort(key=lambda x: RIESGO_ORDER.index(x["riesgo"]) if x["riesgo"] in RIESGO_ORDER else 99)
    return filas


def _load_cripto_cross(g, fx_usd_eur):
    """GUI!D15:H19 -> cruce cartera x nivel de riesgo.

    Estos valores estan en USD sin convertir en la propia hoja GUI (igual
    que GUI!D8:D12); se multiplican aqui por el tipo de cambio para que
    reconcilien con el total oficial de cripto en EUR.
    """
    carteras_cols = {"D": "Cartera Old", "E": "Cartera New", "F": "RD26", "G": "DEFI", "H": "HYPE"}
    filas = []
    for r in range(15, 20):
        riesgo = _norm_riesgo(g.get(f"C{r}"))
        if not riesgo:
            continue
        for col, cartera in carteras_cols.items():
            valor = g.get(f"{col}{r}")
            if isinstance(valor, (int, float)) and valor > 0:
                filas.append({"riesgo": riesgo, "cartera": cartera, "valor": valor * fx_usd_eur})
    return filas


def _load_cripto_uso_cross(g, fx_usd_eur):
    """GUI!T35:AA38 -> cruce cartera x usufructo (Hold/Stake/Farm), por cada
    una de las 5 carteras individuales. La cabecera de esta tabla en el
    propio Excel llama a la 3a cartera 'RD_26' (con guion bajo, unico sitio
    de todo el libro que la nombra asi); se normaliza aqui a 'RD26' para que
    coincida con CARTERA_ORDER. Viene en USD sin convertir (igual que el
    resto de tablas de GUI); se convierte aqui a EUR.
    """
    carteras_cols = {"U": "Cartera Old", "V": "Cartera New", "W": "RD26", "X": "DEFI", "Y": "HYPE"}
    usos_filas = {36: "Hold", 37: "Stake", 38: "Farm"}
    filas = []
    for r, uso in usos_filas.items():
        for col, cartera in carteras_cols.items():
            valor = g.get(f"{col}{r}")
            if isinstance(valor, (int, float)) and valor > 0:
                filas.append({"uso": uso, "cartera": cartera, "valor": valor * fx_usd_eur})
    return filas


def _load_cripto_posiciones(g, fx_usd_eur):
    """Portfolio!M9:Q164 -> una fila por posicion cripto individual.

    O (Capital) y Q (Precio) estan en USD en la hoja; se convierten aqui a
    EUR con el tipo de cambio de Portfolio!D3 para que sumen exactamente el
    total de cripto de CashFlow/GUI (en EUR).
    """
    filas = []
    for r in range(9, 165):
        riesgo = _norm_riesgo(g.get(f"M{r}"))
        coin = g.get(f"N{r}")
        if riesgo is None or coin is None:
            continue
        capital_usd = g.get(f"O{r}") or 0
        cantidad = g.get(f"P{r}")
        precio_usd = g.get(f"Q{r}")
        filas.append({
            "coin": coin,
            "riesgo": riesgo,
            "capital": capital_usd * fx_usd_eur,
            "cantidad": cantidad,
            "precio": precio_usd * fx_usd_eur if isinstance(precio_usd, (int, float)) else None,
        })
    return filas


def _load_cripto_top_por_cartera(g, fx_usd_eur):
    """Portfolio!N9:N164 + columna de capital propia de cada cartera (ver
    CARTERA_CAPITAL_COL) -> para cada una de las 5 carteras, sus posiciones
    individuales (coin, capital) con capital > 0 — usado para el top 10 de
    'Detalle por cartera'. Capital en USD sin convertir; se convierte aqui a
    EUR.
    """
    resultado = {cartera: [] for cartera in CARTERA_CAPITAL_COL}
    for r in range(9, 165):
        coin = g.get(f"N{r}")
        if coin is None:
            continue
        for cartera, col_cap in CARTERA_CAPITAL_COL.items():
            capital_usd = g.get(f"{col_cap}{r}")
            if isinstance(capital_usd, (int, float)) and capital_usd > 0:
                resultado[cartera].append({"coin": coin, "capital": capital_usd * fx_usd_eur})
    return resultado


def _fx_semanal_por_fila(g, filas_fecha):
    """Evolucion!BD -> tipo de cambio EUR/USD de cada fila semanal, con
    relleno hacia adelante y, para las filas anteriores al primer tipo de
    cambio conocido, relleno hacia atras con ese mismo primer valor (BD solo
    se registra a partir de mayo de 2025 — ver _load_cartera_evolucion).
    """
    fx_por_fila = {}
    ultimo_fx = None
    for r, _ in filas_fecha:
        bd = g.get(f"BD{r}")
        if isinstance(bd, (int, float)):
            ultimo_fx = bd
        fx_por_fila[r] = ultimo_fx
    primer_fx = next((v for v in fx_por_fila.values() if v is not None), 1.0)
    for r in fx_por_fila:
        if fx_por_fila[r] is None:
            fx_por_fila[r] = primer_fx
    return fx_por_fila


def _load_cartera_evolucion(g, filas_fecha, fx_por_fila):
    """Evolucion!D205:AJ~317 -> serie semanal de valor actual y capital
    invertido para cada una de las 5 carteras cripto (ver
    CARTERA_EVOLUCION_COLS). Devuelve un dict {cartera: [filas]}.

    Fila 205 es donde arrancan las 5 series (verificado: 0 filas con dato
    antes de esa fila para cualquiera de las 5); de ahi en adelante se para
    dinamicamente (ver _iter_fechas), no en una fila fija.

    Las columnas X:AJ estan en USD sin convertir (igual que Portfolio!O/Q):
    verificado comparando (suma de las 5 carteras * Evolucion!BD de esa
    fila) contra el total oficial semanal de cripto en EUR (Evolucion!I),
    que reconcilia al centimo desde que existe esa columna. BD (tipo de
    cambio EUR/USD semanal) solo se registra a partir de mayo de 2025; para
    las semanas anteriores no hay tipo de cambio propio de esa fecha, asi
    que se usa el primer tipo de cambio conocido como aproximacion (mantiene
    la escala en EUR de toda la serie en vez de mezclar USD y EUR).
    """
    resultado = {cartera: [] for cartera in CARTERA_EVOLUCION_COLS}
    for r, fecha in filas_fecha:
        fx = fx_por_fila[r]
        for cartera, (col_actual, col_inv) in CARTERA_EVOLUCION_COLS.items():
            actual = g.get(f"{col_actual}{r}")
            if not isinstance(actual, (int, float)):
                continue
            invertido = g.get(f"{col_inv}{r}") if col_inv else None
            invertido = invertido if isinstance(invertido, (int, float)) else None
            resultado[cartera].append({
                "fecha": fecha,
                "actual": actual * fx,
                "invertido": invertido * fx if invertido is not None else None,
            })
    return resultado


def _load_cripto_riesgo_evolucion(g, filas_fecha, fx_por_fila):
    """Evolucion!AL:AP -> serie semanal de valor por nivel de riesgo
    (Liquided/LP/Medium/Gem/Crash) del conjunto de las 5 carteras, en EUR.
    Devuelve un dict {riesgo: [filas]}. Mismo patron y misma conversion de
    divisa que _load_cartera_evolucion (USD sin convertir en la hoja).
    """
    resultado = {tier: [] for tier in RIESGO_EVOLUCION_COLS}
    for r, fecha in filas_fecha:
        fx = fx_por_fila[r]
        for tier, col in RIESGO_EVOLUCION_COLS.items():
            valor = g.get(f"{col}{r}")
            if not isinstance(valor, (int, float)):
                continue
            resultado[tier].append({"fecha": fecha, "valor": valor * fx})
    return resultado


def _load_cripto_uso_evolucion(g, filas_fecha, fx_por_fila):
    """Evolucion!AQ:AS -> serie semanal de valor por usufructo (Hold/Stake/
    Farm) del conjunto de las 5 carteras, en EUR. Devuelve un dict
    {uso: [filas]}. Mismo patron que _load_cripto_riesgo_evolucion.
    """
    resultado = {tier: [] for tier in USO_EVOLUCION_COLS}
    for r, fecha in filas_fecha:
        fx = fx_por_fila[r]
        for tier, col in USO_EVOLUCION_COLS.items():
            valor = g.get(f"{col}{r}")
            if not isinstance(valor, (int, float)):
                continue
            resultado[tier].append({"fecha": fecha, "valor": valor * fx})
    return resultado


def _load_cripto_rendimiento(g, fx_usd_eur):
    """GUI!C8:K12 -> tabla de rendimiento por cartera: actual, invertido,
    profit y variacion en 5 ventanas temporales (1 semana, 1/3/6 meses, 1 ano).

    Actual/Invertido vienen en USD sin convertir en la propia hoja (igual que
    GUI!D8:D12, ver _load_cripto_cross): se convierten aqui a EUR — se
    verifico que Actual*fx reconcilia exactamente con el valor oficial de
    Portfolio!H12:K16 (cripto_cartera). Profit y las variaciones +/- ya son
    fracciones (0.23 = +23%) y no necesitan conversion de divisa. Una celda
    con error de Excel (p.ej. '#DIV/0!' cuando Invertido es 0) se guarda como
    None.
    """
    def _num(v):
        return v if isinstance(v, (int, float)) else None

    filas = []
    for r in range(8, 13):
        nombre = g.get(f"C{r}")
        if not nombre:
            continue
        actual = _num(g.get(f"D{r}"))
        invertido = _num(g.get(f"E{r}"))
        filas.append({
            "cartera": nombre,
            "actual": actual * fx_usd_eur if actual is not None else None,
            "invertido": invertido * fx_usd_eur if invertido is not None else None,
            "profit_pct": _num(g.get(f"F{r}")),
            "ret_1s": _num(g.get(f"G{r}")),
            "ret_1m": _num(g.get(f"H{r}")),
            "ret_3m": _num(g.get(f"I{r}")),
            "ret_6m": _num(g.get(f"J{r}")),
            "ret_1a": _num(g.get(f"K{r}")),
        })
    filas.sort(key=lambda x: CARTERA_ORDER.index(x["cartera"]) if x["cartera"] in CARTERA_ORDER else 99)
    return filas


def _load_rv_posiciones(g):
    """Portfolio!A35:K58 -> una fila por posicion de renta variable / MP / liquidez broker."""
    filas = []
    for r in range(35, 60):
        campo = g.get(f"H{r}")
        activo = g.get(f"I{r}")
        total = g.get(f"J{r}")
        if campo is None or activo is None or total is None:
            continue
        k_val = g.get(f"K{r}")
        filas.append({
            "categoria": g.get(f"A{r}"),
            "campo": campo,
            "activo": activo,
            "invertido": g.get(f"F{r}"),
            "valor_actual": total,
            "rendimiento_pct": k_val * 100 if isinstance(k_val, (int, float)) else None,
        })
    return filas


def aggregate_bancos(bancos):
    """Agrupa las filas de bancos por 'proveedor' (columna C de Portfolio,
    o la descripcion cuando C viene vacia) para las graficas de composicion.
    """
    agg = {}
    for b in bancos:
        key = b["proveedor"]
        agg[key] = agg.get(key, 0) + b["importe"]
    return [{"proveedor": k, "importe": v} for k, v in agg.items()]


def _load_rv_resumen(g):
    """GUI!C91:D92 -> capital invertido TOTAL oficial de renta variable.

    OJO: NO se usa la suma de 'invertido' por posicion (Portfolio!F35:F58)
    como total, porque esa suma (~6.512€) NO coincide con esta celda
    dedicada 'RV_Inv' (~7.928€) que el propio Excel usa para calcular su
    rendimiento global de RV (GUI!F91 = 1,95%, coincide con estos dos
    numeros). La celda 'RV_Inv' es la que el usuario trata como oficial,
    asi que es la que mostramos como KPI; el detalle por posicion sigue
    usando el invertido individual de Portfolio, que es mas fino pero no
    suma exactamente igual (puede haber posiciones ya vendidas u otros
    ajustes que 'RV_Inv' arrastra y que no estan en la tabla de posiciones
    actuales).
    """
    actual = None
    invertido = None
    for r in range(1, 93):
        label = g.get(f"C{r}")
        if label is None:
            continue
        label_up = str(label).strip().upper()
        if label_up == "RV":
            actual = g.get(f"D{r}")
        elif label_up == "RV_INV":
            invertido = g.get(f"D{r}")
    rendimiento_pct = None
    if isinstance(actual, (int, float)) and isinstance(invertido, (int, float)) and invertido:
        rendimiento_pct = (actual - invertido) / invertido * 100
    return {"actual": actual, "invertido": invertido, "rendimiento_pct": rendimiento_pct}


def _load_evolucion(g):
    """Evolucion!D3:D~317 (fecha), F (BANCA/liquidez) y K (Patrimonio Activo:
    cripto + renta variable + otros activos) -> serie semanal de patrimonio
    total apilado en 2 capas, igual que 'Grafico 1' de la hoja CashFlow del
    propio Excel (area apilada). Para dinamicamente (ver _iter_fechas), no
    en una fila fija.
    """
    filas = []
    for r, fecha in _iter_fechas(g, start_row=3):
        banca = g.get(f"F{r}")
        activos = g.get(f"K{r}")
        if not isinstance(banca, (int, float)) or not isinstance(activos, (int, float)):
            continue
        filas.append({"fecha": fecha, "banca": banca, "activos": activos})
    return filas


def _load_rv_evolucion(g):
    """Evolucion!D286:D~317 (fecha), Q (RV: valor actual) y R (RV_inv: capital
    invertido) -> serie semanal reciente, igual que 'Grafico 1' de la hoja
    GUI del propio Excel (linea, no area apilada). Rango mas corto que
    _load_evolucion: son solo las semanas recientes que registra esa parte
    de la hoja. Para dinamicamente (ver _iter_fechas), no en una fila fija.
    """
    filas = []
    for r, fecha in _iter_fechas(g, start_row=286):
        rv_actual = g.get(f"Q{r}")
        rv_inv = g.get(f"R{r}")
        if not isinstance(rv_actual, (int, float)) or not isinstance(rv_inv, (int, float)):
            continue
        filas.append({"fecha": fecha, "actual": rv_actual, "invertido": rv_inv})
    return filas


def _load_rv_categoria_evolucion(g):
    """Evolucion!S:W -> serie semanal de valor por categoria de RV
    (Liquided/Fondos/MP/ETF/Acciones). Devuelve un dict {categoria: [filas]},
    mismo rango de fechas que _load_rv_evolucion (Evolucion!D286:D~317).
    """
    resultado = {cat: [] for cat in RV_CATEGORIA_EVOLUCION_COLS}
    for r, fecha in _iter_fechas(g, start_row=286):
        for cat, col in RV_CATEGORIA_EVOLUCION_COLS.items():
            valor = g.get(f"{col}{r}")
            if not isinstance(valor, (int, float)):
                continue
            resultado[cat].append({"fecha": fecha, "valor": valor})
    return resultado


def _load_rv_categoria(g):
    """Portfolio!C60:E64 -> desglose RV por tipo de producto (Liquidez/Fondos/MP/ETF/Acciones)."""
    filas = []
    for r in range(60, 65):
        cat = g.get(f"C{r}")
        importe = g.get(f"D{r}")
        if cat is None or importe is None:
            continue
        filas.append({"categoria": cat, "importe": importe})
    return filas


def load_data_checked(max_attempts=3, delay_seconds=0.5):
    """Como load_data(), pero reintenta si detecta que las cifras no
    reconcilian (probable lectura a medio-sincronizar de finance.xlsx en
    Google Drive). Devuelve el mejor intento (el ultimo) con su lista de
    'warnings' incluida, aunque siga sin cuadrar tras los reintentos.
    """
    import time
    result = None
    for attempt in range(max_attempts):
        result = load_data()
        if not result["warnings"]:
            return result
        if attempt < max_attempts - 1:
            time.sleep(delay_seconds)
    return result


if __name__ == "__main__":
    import json
    import time
    t0 = time.time()
    d = load_data_checked()
    print(f"tiempo de carga: {time.time() - t0:.2f}s")
    warnings = d.pop("warnings")
    print(json.dumps({k: (v if not isinstance(v, list) else f"{len(v)} filas") for k, v in d.items()}, indent=2, ensure_ascii=False, default=str))
    print()
    print("warnings:", warnings if warnings else "(ninguno, todo reconcilia)")
